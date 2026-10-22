"""Generic generated fixtures only; no customer data or privacy certification."""
import io
import json
from pathlib import Path
import sys
import unittest
from unittest.mock import patch
import zipfile

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'skills/data-model-accelerator/scripts'))
import sensitive_data as scanner
from test_sensitive_data_ooxml import archive, replace_member

try:
    from docx import Document
    from PIL import Image, PngImagePlugin
    import pypdf
    from pypdf import PdfWriter
    from pypdf.generic import (ArrayObject, DecodedStreamObject, DictionaryObject,
                              NameObject, TextStringObject)
    AVAILABLE = True
except ImportError:
    AVAILABLE = False


def text_pdf(text='Synthetic migration plan', mutate=None):
    writer = PdfWriter()
    page = writer.add_blank_page(612, 792)
    font = DictionaryObject({NameObject('/Type'): NameObject('/Font'),
                             NameObject('/Subtype'): NameObject('/Type1'),
                             NameObject('/BaseFont'): NameObject('/Helvetica')})
    page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({
        NameObject('/F1'): writer._add_object(font)})})
    contents = DecodedStreamObject()
    contents.set_data(('BT /F1 12 Tf 72 720 Td (' + text + ') Tj ET').encode())
    page[NameObject('/Contents')] = writer._add_object(contents)
    if mutate:
        mutate(writer, page)
    out = io.BytesIO(); writer.write(out)
    return out.getvalue()


@unittest.skipUnless(AVAILABLE, 'Install requirements-dev.txt for qualified rich-document fixtures')
class RichDocumentScanTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        image = io.BytesIO()
        Image.new('RGB', (80, 40), 'white').save(image, format='PNG')
        cls.png = image.getvalue()
        document = Document()
        document.add_paragraph('Synthetic migration plan')
        document.add_picture(io.BytesIO(cls.png))
        out = io.BytesIO(); document.save(out)
        cls.document = out.getvalue()

    def assert_unsupported(self, body, name, reason=None):
        report = scanner.scan_bytes(body, name)
        self.assertEqual(report['status'], 'unsupported', report)
        self.assertFalse(report['coverage']['complete'])
        if reason:
            self.assertIn(reason, {row['rule_id'] for row in report['findings']})
        return report

    def test_generated_docx_png_and_default_jpeg_thumbnail_nested_archive(self):
        with zipfile.ZipFile(io.BytesIO(self.document)) as doc:
            self.assertIn('word/media/image1.png', doc.namelist())
            self.assertIn('docProps/thumbnail.jpeg', doc.namelist())
        for body, name in ((self.document, 'plan.docx'),
                           (archive([('plan.docx', self.document)]), 'delivery.zip')):
            report = scanner.scan_bytes(body, name)
            self.assertEqual(report['status'], 'clear', report)
            self.assertTrue(report['coverage']['complete'])
            self.assertTrue(any('not OCR-scanned' in limitation for limitation in report['limitations']))
        # Raster support is restricted to recognized Office media locations.
        self.assert_unsupported(self.png, 'image.png')
        self.assert_unsupported(archive([('image.png', self.png)]), 'arbitrary.zip')

    def test_image_metadata_is_scanned_without_disclosing_values(self):
        info = PngImagePlugin.PngInfo()
        info.add_text('Description', 'SYNTHETIC_PII_CANARY', zip=True)
        image = io.BytesIO()
        Image.new('RGB', (80, 40), 'white').save(image, format='PNG', pnginfo=info)
        content = replace_member(self.document, 'word/media/image1.png', image.getvalue())
        report = scanner.scan_bytes(content, 'plan.docx')
        self.assertEqual(report['status'], 'blocked')
        self.assertNotIn('SYNTHETIC_PII_CANARY', json.dumps(report))

    def test_malformed_images_documents_and_extra_payloads_fail_closed(self):
        mutations = [('word/media/image1.png', self.png[:-2]),
                     ('word/media/image1.png', self.png + b'MZ arbitrary executable'),
                     ('word/media/image1.png', b'MZ arbitrary binary'),
                     ('docProps/thumbnail.jpeg', b'\xff\xd8invalid\xff\xd9'),
                     ('word/document.xml', b'<w:document><broken>')]
        for name, value in mutations:
            with self.subTest(member=name):
                self.assert_unsupported(replace_member(self.document, name, value), 'plan.docx')
        self.assert_unsupported(archive([('word/media/image1.png', self.png)]), 'plan.docx',
                                'coverage.invalid_office_package')
        with patch.object(scanner, 'MAX_IMAGE_PIXELS', 10):
            self.assert_unsupported(self.document, 'plan.docx', 'coverage.image_pixel_limit')

    def test_text_pdf_and_nested_delivery_scan_clear(self):
        pdf = text_pdf()
        for body, name in ((pdf, 'plan.pdf'), (archive([('plan.pdf', pdf)]), 'delivery.zip')):
            report = scanner.scan_bytes(body, name)
            self.assertEqual(report['status'], 'clear', report)
            self.assertTrue(report['coverage']['complete'])

    def test_pdf_visible_text_metadata_and_unreferenced_strings_are_value_free(self):
        value = 'SYNTHETIC_PII_CANARY'
        bodies = [text_pdf(value), text_pdf(mutate=lambda w, p: w.add_metadata({'/Author': value})),
                  text_pdf(mutate=lambda w, p: w._add_object(TextStringObject(value)))]
        for body in bodies:
            report = scanner.scan_bytes(body, 'plan.pdf')
            self.assertEqual(report['status'], 'blocked', report)
            self.assertIn('pii.synthetic_marker', {row['rule_id'] for row in report['findings']})
            self.assertNotIn(value, json.dumps(report))

    def test_pdf_comments_are_scanned_and_inline_images_fail_closed(self):
        pdf = text_pdf()
        # Comment immediately before EOF does not alter xref object offsets.
        pdf = pdf.replace(b'%%EOF', b'% SYNTHETIC_PII_CANARY\n%%EOF')
        report = scanner.scan_bytes(pdf, 'plan.pdf')
        self.assertEqual(report['status'], 'blocked')
        self.assertNotIn('SYNTHETIC_PII_CANARY', json.dumps(report))
        def inline(writer, page):
            contents = page['/Contents']
            contents.set_data(contents.get_data() + b' q BI /W 1 /H 1 /CS /RGB /BPC 8 ID \xff\x00\x00 EI Q\n')
        self.assert_unsupported(text_pdf(mutate=inline), 'plan.pdf', 'coverage.pdf_inline_image')

    def test_tagged_pdf_indirect_chains_and_passive_layout_attributes(self):
        def tagged(writer, page):
            previous = None
            for _ in range(40):
                element = DictionaryObject({NameObject('/Type'): NameObject('/StructElem'),
                    NameObject('/S'): NameObject('/P'), NameObject('/A'): DictionaryObject({
                        NameObject('/O'): NameObject('/Layout'), NameObject('/Placement'): NameObject('/Block')})})
                if previous is not None:
                    element[NameObject('/K')] = previous
                previous = writer._add_object(element)
            writer._root_object[NameObject('/StructTreeRoot')] = writer._add_object(DictionaryObject({
                NameObject('/Type'): NameObject('/StructTreeRoot'), NameObject('/K'): previous}))
        report = scanner.scan_bytes(text_pdf(mutate=tagged), 'plan.pdf')
        self.assertEqual(report['status'], 'clear', report)

    def test_structure_attributes_do_not_allow_action_dictionaries(self):
        def malicious(writer, page):
            writer._add_object(DictionaryObject({NameObject('/Type'): NameObject('/StructElem'),
                NameObject('/A'): DictionaryObject({NameObject('/O'): NameObject('/Layout'),
                    NameObject('/S'): NameObject('/JavaScript')})}))
        self.assert_unsupported(text_pdf(mutate=malicious), 'plan.pdf',
                                'coverage.pdf_active_or_external_content')
        def action(writer, page):
            page[NameObject('/A')] = DictionaryObject({NameObject('/S'): NameObject('/GoTo')})
        self.assert_unsupported(text_pdf(mutate=action), 'plan.pdf',
                                'coverage.pdf_active_or_external_content')

    def test_pdf_active_content_attachments_external_actions_and_xfa_rejected(self):
        cases = [lambda w, p: w.add_js('app.alert("synthetic")'),
                 lambda w, p: w.add_attachment('synthetic.txt', b'harmless'),
                 lambda w, p: w._root_object.update({NameObject('/XFA'): TextStringObject('synthetic')}),
                 lambda w, p: w._add_object(DictionaryObject({NameObject('/S'): NameObject('/URI'),
                     NameObject('/URI'): TextStringObject('https://example.invalid/')})),
                 lambda w, p: p.update({NameObject('/AA'): DictionaryObject()})]
        for mutate in cases:
            self.assert_unsupported(text_pdf(mutate=mutate), 'plan.pdf',
                                    'coverage.pdf_active_or_external_content')
        self.assert_unsupported(text_pdf(mutate=lambda w, p: w.encrypt('synthetic')), 'plan.pdf',
                                'coverage.pdf_encrypted')

    def test_pdf_image_only_malformed_unknown_stream_and_decoded_budget_rejected(self):
        self.assert_unsupported(text_pdf(''), 'plan.pdf',
                                'coverage.pdf_image_only_or_unextractable_page')
        for body in (b'%PDF-1.7\nbroken\n%%EOF', text_pdf() + b'MZ extra payload'):
            self.assert_unsupported(body, 'plan.pdf')
        def opaque(writer, page):
            stream = DecodedStreamObject(); stream.set_data(b'opaque binary\0payload')
            writer._add_object(stream)
        self.assert_unsupported(text_pdf(mutate=opaque), 'plan.pdf', 'coverage.pdf_unclassified_stream')
        with patch.object(scanner, 'MAX_TOTAL_BYTES', len(text_pdf()) + 1):
            self.assert_unsupported(text_pdf(), 'plan.pdf', 'coverage.pdf_decoded_byte_limit')

    def test_dependency_qualification_errors_are_actionable_and_value_free(self):
        for version in ('6.10.0', 'unqualified'):
            with self.subTest(reader='pypdf', version=version), patch.object(pypdf, '__version__', version):
                with self.assertRaisesRegex(scanner._UnsupportedDocument, 'pdf_dependency_unqualified'):
                    scanner._pdf_worker(text_pdf(), scanner.MAX_TOTAL_BYTES)
        for version in ('12.3.0', 'unqualified'):
            with self.subTest(reader='Pillow', version=version), patch('PIL.__version__', version):
                self.assert_unsupported(self.document, 'plan.docx', 'coverage.image_dependency_unqualified')


if __name__ == '__main__':
    unittest.main()
